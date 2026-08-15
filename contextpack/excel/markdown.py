"""Render Excel values safely as compact Markdown tables."""

from __future__ import annotations

from typing import Any, Iterable

from openpyxl.utils import get_column_letter


def display(value: Any) -> str:
    """Escape a cell value for use inside a Markdown table."""

    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.replace("|", "\\|").replace("\n", "<br>")


def rectangular_table(worksheet: Any, max_row: int, max_col: int) -> str:
    """Render a reasonably sized used range as a conventional table."""

    rows = ["| Row | " + " | ".join(get_column_letter(i) for i in range(1, max_col + 1)) + " |"]
    rows.append("| ---: | " + " | ".join("---" for _ in range(max_col)) + " |")
    for row_idx in range(1, max_row + 1):
        values = [display(worksheet.cell(row_idx, col_idx).value) for col_idx in range(1, max_col + 1)]
        rows.append(f"| {row_idx} | " + " | ".join(values) + " |")
    return "\n".join(rows) + "\n"


def sparse_table(cells: Iterable[Any]) -> str:
    """Render only populated coordinates when the declared range is enormous."""

    rows = ["| Cell | Value |", "| --- | --- |"]
    for cell in sorted(cells, key=lambda item: (item.row, item.column)):
        rows.append(f"| {cell.coordinate} | {display(cell.value)} |")
    return "\n".join(rows) + "\n"

