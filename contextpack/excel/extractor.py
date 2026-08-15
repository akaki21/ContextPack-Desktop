"""Orchestrate safe extraction of an Excel workbook into package reports."""

from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .analysis import analyze_sheet
from .reporting import WorkbookReport
from .workbook import calculation_mode, safe_sheet_folder


@dataclass(frozen=True)
class ExtractionResult:
    """Small completion summary used by command-line and GUI progress output."""

    sheet_count: int
    formula_count: int


def extract_workbook(source: Path, output: Path) -> ExtractionResult:
    """Extract one workbook and guarantee that every opened handle is closed."""

    source = source.resolve()
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output / source.name)

    keep_vba = source.suffix.lower() in {".xlsm", ".xltm"}
    formulas_book: Any | None = None
    values_book: Any | None = None
    try:
        formulas_book = load_workbook(
            source,
            data_only=False,
            read_only=False,
            keep_vba=keep_vba,
            keep_links=True,
        )
        values_book = load_workbook(
            source,
            data_only=True,
            read_only=False,
            keep_vba=keep_vba,
            keep_links=True,
        )
        sheets_root = output / "sheets-data"
        sheets_root.mkdir(parents=True, exist_ok=True)

        external_links = len(getattr(formulas_book, "_external_links", []))
        report = WorkbookReport(source.name, external_links)
        for index, formula_worksheet in enumerate(formulas_book.worksheets, 1):
            value_worksheet = values_book[formula_worksheet.title]
            analysis = analyze_sheet(formula_worksheet, value_worksheet)
            folder_name = safe_sheet_folder(index, formula_worksheet.title)
            report.add_sheet(
                index=index,
                worksheet=formula_worksheet,
                folder_name=folder_name,
                sheet_directory=sheets_root / folder_name,
                analysis=analysis,
            )

        report.write_summary(
            output,
            defined_names=len(formulas_book.defined_names),
            calculation_mode=calculation_mode(formulas_book),
        )
        return ExtractionResult(len(report.sheet_metrics), report.total_formulas)
    finally:
        # The first workbook must also close if opening or processing the second
        # one fails. This prevents locked source files in long-running GUI jobs.
        if values_book is not None:
            values_book.close()
        if formulas_book is not None:
            formulas_book.close()

