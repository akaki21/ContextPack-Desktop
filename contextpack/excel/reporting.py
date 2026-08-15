"""Write the stable Markdown and JSON files in an Excel context package."""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter

from .analysis import SheetAnalysis
from .markdown import display


@dataclass
class WorkbookReport:
    """Accumulate worksheet results and write the workbook-level reports."""

    source_name: str
    external_links: int
    values_index: list[str] = field(init=False)
    formulas_index: list[str] = field(init=False)
    info_parts: list[str] = field(init=False)
    quality_warnings: list[str] = field(default_factory=list)
    sheet_metrics: list[dict[str, Any]] = field(default_factory=list)
    total_formulas: int = 0
    total_errors: int = 0

    def __post_init__(self) -> None:
        self.values_index = [f"# Displayed values index — {self.source_name}", ""]
        self.formulas_index = [f"# Formulas index — {self.source_name}", ""]
        self.info_parts = [f"# Workbook information — {self.source_name}", ""]
        if self.external_links:
            self.quality_warnings.append(f"Workbook contains {self.external_links} external link(s).")

    def add_sheet(
        self,
        *,
        index: int,
        worksheet: Any,
        folder_name: str,
        sheet_directory: Path,
        analysis: SheetAnalysis,
    ) -> None:
        """Write one sheet's files and add its indexes, summary, and metrics."""

        sheet_directory.mkdir(parents=True, exist_ok=True)
        if analysis.warning:
            self.quality_warnings.append(analysis.warning)
        self.total_formulas += len(analysis.formulas)
        self.total_errors += len(analysis.cached_errors)

        (sheet_directory / "values.md").write_text(
            f"# Values — {worksheet.title}\n\n{analysis.values_markdown}",
            encoding="utf-8",
        )
        formula_lines = [f"# Formulas — {worksheet.title}", ""]
        if analysis.formulas:
            formula_lines.extend(
                [
                    "| Cell | Formula | Cached result | Number format |",
                    "| --- | --- | --- | --- |",
                ]
            )
            formula_lines.extend(
                f"| {record.coordinate} | {display(record.formula)} | {display(record.cached_result)} | {display(record.number_format)} |"
                for record in analysis.formulas
            )
        else:
            formula_lines.append("_No formulas._")
        (sheet_directory / "formulas.md").write_text("\n".join(formula_lines) + "\n", encoding="utf-8")

        relative_values = f"sheets-data/{folder_name}/values.md"
        relative_formulas = f"sheets-data/{folder_name}/formulas.md"
        self.values_index.append(f"- [{index}. {worksheet.title}]({relative_values})")
        self.formulas_index.append(
            f"- [{index}. {worksheet.title}]({relative_formulas}) — {len(analysis.formulas)} formula(s)"
        )
        self.info_parts.extend(
            [
                f"## {index}. {worksheet.title}",
                f"- Visibility: {worksheet.sheet_state}",
                f"- Populated cells: {analysis.populated_cells}",
                f"- Populated bounds: A1:{get_column_letter(analysis.max_column)}{analysis.max_row}"
                if analysis.max_row and analysis.max_column
                else "- Populated bounds: empty",
                f"- Output mode: {analysis.output_mode}",
                f"- Formulas: {len(analysis.formulas)}",
                f"- Cached formula errors: {len(analysis.cached_errors)}",
                f"- Merged ranges: {analysis.merged_ranges}",
                f"- Hidden rows / columns: {analysis.hidden_rows} / {analysis.hidden_columns}",
                f"- Charts / embedded images: {analysis.charts} / {analysis.images}",
                f"- Values: [{relative_values}]({relative_values})",
                f"- Formulas: [{relative_formulas}]({relative_formulas})",
                "",
            ]
        )
        self.sheet_metrics.append(
            analysis.metrics(index=index, title=worksheet.title, visibility=worksheet.sheet_state)
        )

    def write_summary(self, output: Path, *, defined_names: int, calculation_mode: str) -> None:
        """Write workbook indexes, quality guidance, and renderer metrics."""

        summary = [
            f"- Sheets: {len(self.sheet_metrics)}",
            f"- Defined names: {defined_names}",
            f"- External links: {self.external_links}",
            f"- Calculation mode: {calculation_mode}",
            f"- Total formulas: {self.total_formulas}",
            f"- Cached formula errors: {self.total_errors}",
            "",
        ]
        self.info_parts[2:2] = summary
        (output / "values.md").write_text("\n".join(self.values_index) + "\n", encoding="utf-8")
        (output / "formulas.md").write_text("\n".join(self.formulas_index) + "\n", encoding="utf-8")
        (output / "workbook-info.md").write_text("\n".join(self.info_parts), encoding="utf-8")

        quality = [
            f"# Quality report — {self.source_name}",
            "",
            f"- Sheets analyzed: {len(self.sheet_metrics)}",
            f"- Total formulas: {self.total_formulas}",
            f"- Cached formula errors: {self.total_errors}",
            f"- External links: {self.external_links}",
            "",
            "## Warnings",
            "",
        ]
        quality.extend(f"- {warning}" for warning in self.quality_warnings)
        if not self.quality_warnings:
            quality.append("- No structural warnings detected.")
        quality.extend(
            [
                "",
                "Cached values may be stale if Excel did not recalculate and save the workbook before packaging.",
                "",
            ]
        )
        (output / "quality-report.md").write_text("\n".join(quality), encoding="utf-8")
        (output / "excel-metrics.json").write_text(
            json.dumps(
                {
                    "sheet_count": len(self.sheet_metrics),
                    "total_formulas": self.total_formulas,
                    "cached_formula_errors": self.total_errors,
                    "external_links": self.external_links,
                    "warnings": self.quality_warnings,
                    "sheets": self.sheet_metrics,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

